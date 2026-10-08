"""Runs the real tool against local stand-ins for Teal and LinkedIn.

The Teal stand-in (tests/fake_teal/) mirrors the real screens: the contact
list with "+ Add a New Contact" (form in an iframe, Save returns to the list),
and contact pages at /contact-tracker/<id> with dropdown menus, calendar-only
dates that save by themselves, Contact Information and an Edit form.

Needs Playwright and a Chrome/Chromium binary. Set TEAL_RPA_TEST_CHROME to the
browser's path if it isn't found automatically; otherwise these tests skip.
"""

import json
import os
import socket
import threading
import uuid
from functools import partial
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

import openpyxl
import pytest

pytest.importorskip("playwright.sync_api")

from teal_rpa import browser  # noqa: E402
from teal_rpa.linkedin import LinkedInPage  # noqa: E402
from teal_rpa.runner import Options, Prompter, Run  # noqa: E402
from teal_rpa.spreadsheet import load_contacts  # noqa: E402
from teal_rpa.state import DONE, FAILED, NEEDS_ATTENTION, RunState  # noqa: E402
from teal_rpa.teal import TealPage  # noqa: E402

FAKE_SITE = Path(__file__).parent / "fake_teal"

# Teal's export layout, repeated columns included.
HEADERS = ["First Name", "Last Name", "Email Address", "Phone", "Position", "Company", "URL",
           "follow_up_at", "last_contacted_at",
           "contact_intention_type", "contact_relationship_type", "contact_next_step_type",
           "contact_intention_type", "contact_relationship_type", "contact_next_step_type"]


def choice(name):
    return json.dumps({"name": name, "id": "1"}) if name else None


def row(first, last, slug, email=None, phone=None, title="CTO", company="Acme",
        goal=None, relationship=None, status=None, follow_up=None, last_contacted=None):
    picks = [choice(goal), choice(relationship), choice(status)]
    return [first, last, email, phone, title, company, f"https://www.linkedin.com/in/{slug}",
            follow_up, last_contacted, *picks, *picks]


def make_xlsx(path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Upload 1"
    ws.append(HEADERS)
    for r in rows:
        ws.append(r)
    wb.save(path)
    return path


def _chrome_path():
    for path in (os.environ.get("TEAL_RPA_TEST_CHROME"), "/opt/pw-browsers/chromium"):
        if path and os.path.exists(path):
            return path
    try:
        return browser.find_chrome()
    except browser.BrowserError:
        return None


CHROME = _chrome_path()
pytestmark = pytest.mark.skipif(CHROME is None, reason="no Chrome/Chromium available")


class FakeSites:
    def __init__(self):
        self.reset()

    def reset(self):
        self.contacts = {}          # Teal: id -> contact
        self.logged_in = True
        self.json_api = True        # False: Teal's replies aren't JSON (no data for the tool)
        self.ignore_fields = set()  # Teal: changes to these fields don't stick
        self.config = {}            # front-end variations (nativeSelect, unlabelledArrows, slowForm)
        self.linkedin_logged_in = True
        self.linkedin_visits = []
        # Profiles where the overlay link lands on the plain profile page (no box).
        self.linkedin_lands_on_profile = set()
        # LinkedIn: slug -> (emails, phones) shown in the Contact info overlay
        self.linkedin_profiles = {
            "edsoohoo": (["ed@lenovo.com"], ["203-555-0100"]),
            "joshreicher": (["josh@usi.com"], ["203-555-0111", "(212) 555-0199"]),
            "private": ([], []),
        }

    def add_contact(self, **fields):
        ident = str(uuid.uuid4())
        contact = {"id": ident, "first_name": "", "last_name": "", "title": "", "company": "",
                   "email": "", "url": "", "phone": "", "twitter": "", "location": "",
                   "relationship": "", "goal": "", "status": "", "follow_up": "",
                   "last_contacted": "", "created": "2026-10-02"}
        contact.update(fields)
        self.contacts[ident] = contact
        return ident

    def by_name(self, first):
        [match] = [c for c in self.contacts.values() if c["first_name"] == first]
        return match


def linkedin_profile_html(slug, emails, phones):
    """The profile page; its "Contact info" link opens the box without reloading."""
    # "</" escaped so the embedded </script> doesn't end this page's script early.
    box = json.dumps(linkedin_html(slug, emails, phones).split("<!--box-->")[1]).replace("</", "<\\/")
    return f"""<!doctype html><html><body><main><h1>{slug}</h1>
<a id="ci" href="/in/{slug}/overlay/contact-info/">Contact info</a></main><div id="box"></div>
<script>document.getElementById("ci").addEventListener("click", (e) => {{
  e.preventDefault(); history.pushState({{}}, "", e.currentTarget.href);
  document.getElementById("box").innerHTML = {box};
  for (const s of document.querySelectorAll("#box script")) eval(s.textContent);
}});</script></body></html>"""


def linkedin_html(slug, emails, phones):
    sections = [f"<section><h3>Profile</h3><a href=\"https://linkedin.com/in/{slug}\">"
                f"linkedin.com/in/{slug}</a></section>"]
    if phones:
        items = "".join(f"<li><span>{p}</span> <span>(Mobile)</span></li>" for p in phones)
        sections.append(f"<section><h3>Phone</h3><ul>{items}</ul></section>")
    sections.append("<section><h3>Address</h3><a href=\"#\">Trumbull, CT</a></section>")
    if emails:
        links = "".join(f"<a href=\"mailto:{e}\">{e}</a><br>" for e in emails)
        sections.append(f"<section><h3>Email</h3>{links}</section>")
    body = json.dumps("".join(sections))
    return f"""<!doctype html><html><body><main>Profile page</main><!--box-->
<div role="dialog" aria-labelledby="t"><h2 id="t">Contact info</h2><div id="c"></div></div>
<script>setTimeout(() => {{ document.getElementById("c").innerHTML = {body}; }}, 300);</script><!--box-->
</body></html>"""


def _handler(sites, *args, **kwargs):
    class Handler(SimpleHTTPRequestHandler):
        def _send(self, status, body=b"", content_type="application/json", headers=()):
            self.send_response(status)
            if body:
                self.send_header("Content-Type", content_type)
                self.send_header("Content-Length", str(len(body)))
            for name, value in headers:
                self.send_header(name, value)
            self.end_headers()
            self.wfile.write(body)

        def _json(self, data):
            kind = "application/json" if sites.json_api else "text/plain"
            self._send(200, json.dumps(data).encode(), kind)

        def do_GET(self):
            path = self.path.split("?")[0]
            if path.startswith("/in/"):
                slug = path.split("/")[2]
                sites.linkedin_visits.append(path)
                if not sites.linkedin_logged_in:
                    return self._send(302, headers=[("Location", "/authwall?trk=x")])
                overlay = path.rstrip("/").endswith("/overlay/contact-info")
                if overlay and slug in sites.linkedin_lands_on_profile:
                    return self._send(302, headers=[("Location", f"/in/{slug}/")])
                if slug in sites.linkedin_profiles:
                    make = linkedin_html if overlay else linkedin_profile_html
                    html = make(slug, *sites.linkedin_profiles[slug])
                else:
                    html = "<html><body><h1>This page doesn’t exist</h1></body></html>"
                return self._send(200, html.encode(), "text/html; charset=utf-8")
            if path.startswith("/authwall"):
                self.path = "/sign-in.html"
            elif path == "/config.js":
                return self._send(200, f"window.FAKE_CONFIG = {json.dumps(sites.config)};".encode(),
                                  "application/javascript")
            elif path == "/api/contacts":
                return self._json({"data": list(sites.contacts.values())})
            elif path.startswith("/contact-tracker"):
                if not sites.logged_in:
                    return self._send(302, headers=[("Location", "/sign-in.html")])
                self.path = "/contact-tracker.html"
            return super().do_GET()

        def _body(self):
            return json.loads(self.rfile.read(int(self.headers["Content-Length"])))

        def do_POST(self):
            ident = sites.add_contact(**self._body())
            self._json({"data": sites.contacts[ident]})

        def do_PATCH(self):
            ident = self.path.rstrip("/").rsplit("/", 1)[-1]
            changes = {k: v for k, v in self._body().items() if k not in sites.ignore_fields}
            sites.contacts[ident].update(changes)
            self._json({"data": sites.contacts[ident]})

        def log_message(self, *a):
            pass

    return Handler(*args, directory=str(FAKE_SITE), **kwargs)


def _free_port():
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        return s.getsockname()[1]


@pytest.fixture(scope="module")
def server():
    sites = FakeSites()
    srv = ThreadingHTTPServer(("127.0.0.1", 0), partial(_handler, sites))
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    sites.origin = f"http://127.0.0.1:{srv.server_address[1]}"
    sites.url = f"{sites.origin}/contact-tracker"
    yield sites
    srv.shutdown()


@pytest.fixture(scope="module")
def chrome_context(server, tmp_path_factory):
    from playwright.sync_api import sync_playwright

    endpoint = browser.ensure_chrome(
        chrome_path=CHROME, profile_dir=tmp_path_factory.mktemp("profile"), port=_free_port(),
        start_url="about:blank", extra_args=["--headless=new", "--no-sandbox", "--disable-gpu"])
    with sync_playwright() as pw:
        b, context = browser.connect(pw, endpoint)
        yield context
        b.close()


@pytest.fixture
def sites(server, chrome_context):
    server.reset()
    yield server


@pytest.fixture
def page(chrome_context):
    page = chrome_context.new_page()
    yield page
    page.close()


class ScriptedPrompter(Prompter):
    def __init__(self, answers=()):
        self.questions = []
        answers = list(answers)

        def fake_input(question):
            self.questions.append(question)
            if not answers:
                raise EOFError
            return answers.pop(0)

        super().__init__(fake_input)


def make_run(tmp_path, sites, page, rows, answers=(), linkedin=True, reuse=False, **options):
    path = tmp_path / "contacts.xlsx"
    if not reuse:
        make_xlsx(path, rows)
    sheet = load_contacts(path)
    options.setdefault("unattended", True)
    return Run(
        sheet=sheet, state=RunState.for_spreadsheet(path),
        teal=TealPage(page, sites.url, timeout_ms=5_000),
        options=Options(linkedin_delay=0, teal_delay=0, backoff=0.1, settle_ms=300, **options),
        prompter=ScriptedPrompter(answers), failures_dir=tmp_path / "failures",
        linkedin=LinkedInPage(page.context, timeout_ms=5_000, origin=sites.origin) if linkedin else None,
        sleep=lambda s: None,
    )


def run_both(run):
    """Both passes, re-reading the spreadsheet in between like the CLI does."""
    contacts = [c for c in run.sheet.contacts if c.is_valid and not run.state.is_done(c.key)]
    assert run.enrich_all(contacts) is not False
    run.sheet = load_contacts(run.sheet.path)
    keys = {c.key for c in contacts}
    run.load_all([c for c in run.sheet.contacts if c.key in keys])
    return run.summary


def sheet_rows(path):
    ws = openpyxl.load_workbook(path).active
    headers = [c.value for c in ws[1]]
    return [dict(zip(headers, (c.value for c in r))) for r in ws.iter_rows(min_row=2)]


# -- the whole flow ------------------------------------------------------------

def test_full_flow_creates_and_completes_contacts(tmp_path, sites, page):
    run = make_run(tmp_path, sites, page, [
        row("Ed", "Soo Hoo", "edsoohoo", title="WW CTO Global Accounts", company="Lenovo",
            goal="Networking", relationship="Mentor", status="Follow up needed",
            follow_up="2027-01-05"),
        row("Joshua", "Reicher", "joshreicher", goal="Networking", relationship="Co-Worker",
            status="To be contacted", follow_up="2026-08-30"),
        row("Pri", "Vate", "private", goal="Request referral", relationship="Friend",
            status="Meeting scheduled"),
    ], answers=["2", "pri@x.com", ""])  # Josh: 2nd phone; Pri: typed email, no phone

    summary = run_both(run)

    assert [c.name for c in summary.done] == ["Ed Soo Hoo", "Joshua Reicher", "Pri Vate"]
    assert len(sites.contacts) == 3
    ed, josh, pri = sites.by_name("Ed"), sites.by_name("Joshua"), sites.by_name("Pri")
    assert (ed["email"], ed["phone"]) == ("ed@lenovo.com", "(203) 555-0100")
    assert (ed["relationship"], ed["goal"], ed["status"], ed["follow_up"]) == \
        ("Mentor", "Networking", "Follow up needed", "2027-01-05")
    assert (josh["phone"], josh["email"]) == ("(212) 555-0199", "josh@usi.com")
    assert josh["relationship"] == "Co-worker"  # Teal's spelling, matched ignoring case
    assert josh["follow_up"] == "2026-08-30"    # calendar stepped backwards
    assert (pri["email"], pri["phone"], pri["status"]) == ("pri@x.com", "", "Meeting scheduled")
    assert pri["follow_up"] == ""

    # Asked: which of Josh's phones, then Pri's email and phone. Nothing else.
    assert len(run.prompter.questions) == 3
    assert "Which phone?" in run.prompter.questions[0]

    rows = sheet_rows(run.sheet.path)
    assert [r["Email Address"] for r in rows] == ["ed@lenovo.com", "josh@usi.com", "pri@x.com"]
    assert [r["Phone"] for r in rows] == ["203-555-0100", "(212) 555-0199", None]
    assert [r["teal_contact_id"] for r in rows] == [ed["id"], josh["id"], pri["id"]]
    assert [r["rpa_status"] for r in rows] == [DONE] * 3
    assert all(run.state.is_done(c.key) for c in run.sheet.contacts)


def test_existing_contact_is_updated_not_duplicated(tmp_path, sites, page):
    ident = sites.add_contact(first_name="Ed", last_name="Soo Hoo", title="CTO", company="Acme",
                              url="https://linkedin.com/in/edsoohoo/", email="old@lenovo.com",
                              relationship="Friend")
    run = make_run(tmp_path, sites, page, [
        row("Ed", "Soo Hoo", "edsoohoo", email="esoohoo@lenovo.com", phone="203-555-0100",
            relationship="Mentor", status="Follow up needed", follow_up="2026-10-16"),
    ])

    summary = run_both(run)

    assert len(summary.done) == 1 and len(sites.contacts) == 1
    ed = sites.contacts[ident]
    assert (ed["email"], ed["phone"]) == ("esoohoo@lenovo.com", "(203) 555-0100")  # via Edit form
    assert (ed["relationship"], ed["status"], ed["follow_up"]) == \
        ("Mentor", "Follow up needed", "2026-10-16")
    assert sheet_rows(run.sheet.path)[0]["teal_contact_id"] == ident
    assert sites.linkedin_visits == []  # email and phone were already in the sheet


def test_resume_uses_saved_teal_id(tmp_path, sites, page):
    # A previous run created the contact, then stopped before finishing it.
    ident = sites.add_contact(first_name="Ed", last_name="Soo Hoo", title="CTO", company="Acme",
                              url="https://www.linkedin.com/in/edsoohoo", email="ed@x.com")
    run = make_run(tmp_path, sites, page, [
        row("Ed", "Soo Hoo", "edsoohoo", email="ed@x.com", phone="1", status="Thank you sent"),
    ], linkedin=False)
    run.state.remember("https://www.linkedin.com/in/edsoohoo", teal_id=ident)

    summary = run_both(run)

    assert len(summary.done) == 1 and len(sites.contacts) == 1
    assert sites.contacts[ident]["status"] == "Thank you sent"


def test_rerun_skips_done_contacts(tmp_path, sites, page):
    rows = [row("Ed", "Soo Hoo", "edsoohoo", email="ed@x.com", phone="1", goal="Networking")]
    run = make_run(tmp_path, sites, page, rows, linkedin=False)
    run_both(run)
    rerun = make_run(tmp_path, sites, page, rows, linkedin=False, reuse=True)
    summary = run_both(rerun)
    assert summary.done == [] and len(sites.contacts) == 1


# -- fallbacks and page variations --------------------------------------------

def test_without_teal_data_uses_the_contact_list(tmp_path, sites, page):
    sites.json_api = False  # the tool can't read Teal's replies; must use the list
    existing = sites.add_contact(first_name="Joshua", last_name="Reicher",
                                 url="https://www.linkedin.com/in/joshreicher")
    sites.add_contact(first_name="Joshua", last_name="Reicher",  # same name, different person
                      url="https://www.linkedin.com/in/another-josh")
    run = make_run(tmp_path, sites, page, [
        row("Ed", "Soo Hoo", "edsoohoo", email="ed@x.com", phone="1", goal="Networking"),
        row("Joshua", "Reicher", "joshreicher", email="j@x.com", phone="2", status="To be contacted"),
    ], linkedin=False)

    summary = run_both(run)

    assert len(summary.done) == 2 and len(sites.contacts) == 3
    assert sites.by_name("Ed")["goal"] == "Networking"
    assert sites.contacts[existing]["status"] == "To be contacted"
    assert [r["teal_contact_id"] for r in sheet_rows(run.sheet.path)] == \
        [sites.by_name("Ed")["id"], existing]


def test_standard_dropdowns_and_unlabelled_calendar_arrows(tmp_path, sites, page):
    sites.config = {"nativeSelect": True, "unlabelledArrows": True}
    run = make_run(tmp_path, sites, page, [
        row("Ed", "Soo Hoo", "edsoohoo", email="ed@x.com", phone="1", relationship="hiring MANAGER",
            follow_up="2026-12-31", last_contacted="2026-09-01"),
    ], linkedin=False)

    summary = run_both(run)

    assert len(summary.done) == 1
    ed = sites.by_name("Ed")
    assert (ed["relationship"], ed["follow_up"], ed["last_contacted"]) == \
        ("Hiring manager", "2026-12-31", "2026-09-01")


# -- things that don't work out -------------------------------------------------

def test_field_that_wont_save_needs_attention(tmp_path, sites, page):
    sites.ignore_fields = {"status"}
    run = make_run(tmp_path, sites, page, [
        row("Ed", "Soo Hoo", "edsoohoo", email="ed@x.com", phone="1", goal="Networking",
            status="Thank you sent"),
    ], linkedin=False, unattended=False, answers=["r", ""])  # try again, then leave it

    summary = run_both(run)

    [(contact, reason)] = summary.attention
    assert "Status: Teal shows '(blank)', sheet has 'Thank you sent'" in reason
    assert sites.by_name("Ed")["goal"] == "Networking"  # the rest was still done
    assert len(run.prompter.questions) == 2
    assert run.state.status(contact.key) == NEEDS_ATTENTION
    assert sheet_rows(run.sheet.path)[0]["rpa_status"] == NEEDS_ATTENTION
    assert (tmp_path / "failures" / "row-2.png").exists()

    # Next run: retried using the saved ID, never a duplicate.
    sites.ignore_fields = set()
    rerun = make_run(tmp_path, sites, page, [], linkedin=False, reuse=True)
    assert len(run_both(rerun).done) == 1 and len(sites.contacts) == 1


def test_value_that_isnt_a_teal_option(tmp_path, sites, page):
    run = make_run(tmp_path, sites, page, [
        row("Ed", "Soo Hoo", "edsoohoo", email="ed@x.com", phone="1", relationship="Mentr"),
    ], linkedin=False)

    [(_, reason)] = run_both(run).attention

    assert "Relationship: 'Mentr' isn't one of Teal's options" in reason


def test_add_form_validation_error_fails_row(tmp_path, sites, page):
    run = make_run(tmp_path, sites, page, [
        row("Bad", "Email", "bad", email="invalid@x", phone="1"),
    ], linkedin=False)

    [(_, reason)] = run_both(run).failed

    assert "Please enter a valid email" in reason and "check Teal before re-running" in reason
    assert run.state.status("https://www.linkedin.com/in/bad") == FAILED


def test_form_that_never_opens_fails_after_retries(tmp_path, sites, page, monkeypatch):
    sites.config = {"slowForm": True}
    run = make_run(tmp_path, sites, page, [
        row("Ed", "Soo Hoo", "edsoohoo", email="ed@x.com", phone="1"),
    ], linkedin=False)
    run.teal.timeout_ms = 1_000  # the form takes 1.5s, so it "never" shows
    attempts = []
    original = run.teal.open_form
    monkeypatch.setattr(run.teal, "open_form", lambda: attempts.append(1) or original())

    [(_, reason)] = run_both(run).failed

    assert len(attempts) == 3  # first try + 2 retries
    assert "didn't appear" in reason and sites.contacts == {}


def test_dry_run_changes_nothing(tmp_path, sites, page):
    ident = sites.add_contact(first_name="Ed", last_name="Soo Hoo",
                              url="https://www.linkedin.com/in/edsoohoo", relationship="Friend")
    run = make_run(tmp_path, sites, page, [
        row("Ed", "Soo Hoo", "edsoohoo", relationship="Mentor"),
        row("Joshua", "Reicher", "joshreicher", email="j@x.com", phone="2"),
    ], answers=[""], dry_run=True)  # Ed's LinkedIn has one phone and one email: no question

    summary = run_both(run)

    assert len(summary.dry_run) == 2
    assert "Relationship" in summary.dry_run[0][1]
    assert list(sites.contacts) == [ident] and sites.contacts[ident]["relationship"] == "Friend"
    assert sheet_rows(run.sheet.path)[0]["Email Address"] is None
    assert not run.state.path.exists()


# -- LinkedIn and sign-in ------------------------------------------------------

def test_nothing_on_linkedin_saves_evidence_and_asks_once(tmp_path, sites, page):
    rows = [row("Pri", "Vate", "private"), row("Gone", "Away", "no-such-person")]
    run = make_run(tmp_path, sites, page, rows, answers=["", "", "", ""])

    run.enrich_all(run.sheet.contacts)

    assert len(run.prompter.questions) == 4  # email + phone, for each
    assert (tmp_path / "failures" / "linkedin-row-2.txt").exists()
    assert "doesn't exist" in (tmp_path / "failures" / "linkedin-row-3.txt").read_text()
    assert run.summary.nothing == ["Pri Vate", "Gone Away"]

    visits = len(sites.linkedin_visits)
    run.enrich_all(load_contacts(run.sheet.path).contacts)  # already handled: no revisit
    assert len(sites.linkedin_visits) == visits


def test_waits_for_sign_in_to_linkedin_and_teal(tmp_path, sites, page):
    sites.logged_in = sites.linkedin_logged_in = False
    run = make_run(tmp_path, sites, page, [row("Ed", "Soo Hoo", "edsoohoo")])
    questions = []

    def sign_in(question):
        questions.append(question)
        if "LinkedIn" in question:
            sites.linkedin_logged_in = True
        else:
            sites.logged_in = True
        return ""

    run.prompter = Prompter(sign_in)
    summary = run_both(run)

    assert "Sign in to LinkedIn" in questions[0] and "Sign in to Teal" in questions[1]
    assert len(summary.done) == 1


def test_input_closing_stops_cleanly(tmp_path, sites, page):
    run = make_run(tmp_path, sites, page, [row("Pri", "Vate", "private")], answers=[])
    assert run.enrich_all(run.sheet.contacts) is False
    assert run.state.status("https://www.linkedin.com/in/private") == "pending"


def test_linkedin_landing_on_profile_opens_contact_info(tmp_path, sites, page):
    sites.linkedin_lands_on_profile = {"joshreicher"}
    run = make_run(tmp_path, sites, page, [row("Joshua", "Reicher", "joshreicher")],
                   answers=["1"])

    run.enrich_all(run.sheet.contacts)

    assert run.summary.found == ["Joshua Reicher: email, phone"]
    assert run.state.recall("https://www.linkedin.com/in/joshreicher", "email") == "josh@usi.com"
    assert sites.linkedin_visits == ["/in/joshreicher/overlay/contact-info/", "/in/joshreicher/"]


def test_linkedin_box_for_someone_else_is_rejected(tmp_path, sites, page, monkeypatch):
    import tests.test_end_to_end as module
    from teal_rpa.linkedin import LinkedInError

    run = make_run(tmp_path, sites, page, [row("Ed", "Soo Hoo", "edsoohoo")])
    assert run.linkedin.lookup("https://www.linkedin.com/in/edsoohoo").email == "ed@lenovo.com"

    real_html = module.linkedin_html
    monkeypatch.setattr(module, "linkedin_html",
                        lambda slug, emails, phones: real_html("someone-else", emails, phones))
    with pytest.raises(LinkedInError, match="is for /in/someone-else"):
        run.linkedin.lookup("https://www.linkedin.com/in/edsoohoo")


def test_linkedin_error_keeps_evidence(tmp_path, sites, page):
    sites.linkedin_profiles["broken"] = ([], [])
    sites.linkedin_lands_on_profile = {"broken"}
    run = make_run(tmp_path, sites, page, [row("Bro", "Ken", "broken")], answers=["", ""])
    run.linkedin.timeout_ms = 1_000

    def no_link(page):  # the profile's Contact info link can't be found either
        pass

    run.linkedin._click_contact_info_link = no_link
    run.enrich_all(run.sheet.contacts)

    evidence = (tmp_path / "failures" / "linkedin-row-2.txt").read_text()
    assert "didn't appear (the tab shows" in evidence and "/in/broken/" in evidence
    assert (tmp_path / "failures" / "linkedin-row-2.png").exists()
