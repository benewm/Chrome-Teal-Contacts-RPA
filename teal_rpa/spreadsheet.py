"""Read the contacts spreadsheet (.xlsx or .csv) and write values back to it."""

from __future__ import annotations

import csv
import json
import re
import time
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from urllib.parse import urlsplit

# Canonical field -> accepted header names (matched case-insensitively).
# Covers both the plain names ("LinkedIn URL", "Title") and Teal's export
# names ("URL", "Position", "Email Address", "contact_relationship_type").
COLUMN_ALIASES: dict[str, list[str]] = {
    "url": ["LinkedIn URL", "URL", "Link", "LinkedIn", "Profile URL"],
    "first_name": ["First Name", "FirstName"],
    "last_name": ["Last Name", "LastName"],
    "title": ["Title", "Position", "Job Title"],
    "company": ["Company", "Company Name"],
    "email": ["Email", "Email Address", "E-mail"],
    "phone": ["Phone", "Phone Number", "Mobile"],
    "location": ["Location"],
    "twitter": ["Twitter", "twitter_handle", "Twitter Handle"],
    "relationship": ["contact_relationship_type", "Relationship"],
    "goal": ["contact_intention_type", "Goal"],
    "status": ["contact_next_step_type", "Status"],
    "follow_up": ["follow_up_at", "Follow up", "Follow Up"],
    "last_contacted": ["last_contacted_at", "Last contacted", "Last Contacted"],
    "teal_id": ["teal_contact_id", "Teal ID"],
    "rpa_status": ["rpa_status"],
    "rpa_notes": ["rpa_notes"],
}
REQUIRED_FIELDS = ["url", "first_name", "last_name", "title", "company"]
DATE_FIELDS = ["follow_up", "last_contacted"]
CHOICE_FIELDS = ["relationship", "goal", "status"]
# Columns the tool adds to the sheet if they're missing: field -> header.
TOOL_COLUMNS = {"teal_id": "teal_contact_id", "rpa_status": "rpa_status", "rpa_notes": "rpa_notes"}

LINKEDIN_PROFILE_RE = re.compile(r"^https://www\.linkedin\.com/in/[^/]+$")


class SpreadsheetError(Exception):
    """The spreadsheet can't be used as-is (bad format, missing columns...)."""


class SpreadsheetLockedError(SpreadsheetError):
    """The file couldn't be written, usually because it's open in Excel."""


@dataclass
class Contact:
    row_number: int  # 1-based row as shown in Excel (header is row 1)
    url: str
    first_name: str
    last_name: str
    title: str
    company: str
    email: str
    phone: str
    location: str = ""
    twitter: str = ""
    relationship: str = ""
    goal: str = ""
    status: str = ""
    follow_up: str = ""       # ISO date (YYYY-MM-DD) or ""
    last_contacted: str = ""  # ISO date (YYYY-MM-DD) or ""
    teal_id: str = ""
    raw: dict[str, str] = field(default_factory=dict)
    problems: list[str] = field(default_factory=list)  # row can't be processed
    warnings: list[str] = field(default_factory=list)  # processed, but worth a look

    @property
    def name(self) -> str:
        return f"{self.first_name} {self.last_name}".strip() or f"row {self.row_number}"

    @property
    def key(self) -> str:
        """Stable id for run tracking: the normalized URL, else the row number."""
        return self.url or f"row:{self.row_number}"

    @property
    def is_valid(self) -> bool:
        return not self.problems


@dataclass
class Sheet:
    path: Path
    sheet_name: str | None
    headers: list[str]
    column_map: dict[str, int]  # canonical field -> 0-based column index
    contacts: list[Contact]
    warnings: list[str]


def normalize_linkedin_url(value: str) -> str:
    """Canonical form: https://www.linkedin.com/in/<slug> (no query, no slash)."""
    value = value.strip()
    if not value:
        return ""
    if not re.match(r"^https?://", value, re.I):
        value = "https://" + value
    parts = urlsplit(value)
    host = parts.netloc.lower()
    if host.endswith("linkedin.com"):
        host = "www.linkedin.com"
    path = parts.path.rstrip("/")
    if path.lower().startswith("/in/"):
        path = "/in/" + path[4:]
    return f"https://{host}{path}"


def choice_name(value: str) -> str:
    """Teal exports dropdown values as {"name":"Mentor","id":"7"}; plain text works too."""
    value = value.strip()
    if value.startswith("{"):
        try:
            data = json.loads(value)
        except ValueError:
            return value
        if isinstance(data, dict):
            return str(data.get("name") or "").strip()
    return value


def parse_date(value: str) -> date | None:
    """Accepts 2026-10-16, 2026-10-16 09:30:00, 2026-10-16T09:30:00Z, 10/16/2026."""
    value = value.strip()
    if not value:
        return None
    match = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", value)
    if match:
        return date(*map(int, match.groups()))
    match = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{4})$", value)
    if match:
        month, day, year = map(int, match.groups())
        return date(year, month, day)
    raise ValueError(f"not a date: {value!r}")


def _cell_to_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        # Phone numbers typed into Excel come back as 5551234567.0
        return str(int(value))
    if isinstance(value, datetime):
        return value.isoformat(sep=" ") if value.time() != datetime.min.time() else value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _map_columns(headers: list[str]) -> tuple[dict[str, int], dict[str, list[int]], list[str]]:
    """Returns (field -> first column, field -> all columns with that header, warnings)."""
    lookup: dict[str, list[int]] = {}
    for index, header in enumerate(headers):
        norm = header.strip().lower()
        if norm:
            lookup.setdefault(norm, []).append(index)

    column_map: dict[str, int] = {}
    all_columns: dict[str, list[int]] = {}
    for canonical, aliases in COLUMN_ALIASES.items():
        for alias in aliases:
            if alias.lower() in lookup:
                all_columns[canonical] = lookup[alias.lower()]
                column_map[canonical] = all_columns[canonical][0]
                break

    missing = [f for f in REQUIRED_FIELDS if f not in column_map]
    if missing:
        wanted = "; ".join(f"{f}: one of {COLUMN_ALIASES[f]}" for f in missing)
        raise SpreadsheetError(
            f"Missing required column(s) -> {wanted}. Found headers: {[h for h in headers if h]}"
        )

    warnings = []
    for norm, indexes in lookup.items():
        if len(indexes) > 1:
            columns = ", ".join(str(i + 1) for i in indexes)
            warnings.append(f"Column '{headers[indexes[0]]}' appears {len(indexes)} times "
                            f"(columns {columns}); using the first non-empty value.")
    return column_map, all_columns, warnings


def _read_rows(path: Path, sheet_name: str | None) -> tuple[list[list], str | None]:
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            if sheet_name and sheet_name not in wb.sheetnames:
                raise SpreadsheetError(f"Sheet '{sheet_name}' not found. Sheets: {wb.sheetnames}")
            ws = wb[sheet_name] if sheet_name else wb.worksheets[0]
            return [list(r) for r in ws.iter_rows(values_only=True)], ws.title
        finally:
            wb.close()
    if suffix == ".csv":
        return _read_csv(path)[0], None
    raise SpreadsheetError(f"Unsupported file type '{path.suffix}'. Use .xlsx or .csv.")


def _read_csv(path: Path) -> tuple[list[list[str]], str]:
    # Excel on Windows saves CSV as UTF-8 with BOM or as cp1252.
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            with path.open(newline="", encoding=encoding) as f:
                return list(csv.reader(f)), encoding
        except UnicodeDecodeError:
            continue
    raise SpreadsheetError(f"Can't decode {path} as UTF-8 or Windows-1252.")


def load_contacts(path: str | Path, sheet_name: str | None = None) -> Sheet:
    path = Path(path)
    if not path.exists():
        raise SpreadsheetError(f"File not found: {path}")

    rows, actual_sheet = _read_rows(path, sheet_name)
    if not rows:
        raise SpreadsheetError(f"{path} is empty.")

    headers = [_cell_to_str(h) for h in rows[0]]
    column_map, all_columns, warnings = _map_columns(headers)

    contacts: list[Contact] = []
    seen_urls: dict[str, int] = {}
    for offset, row in enumerate(rows[1:]):
        row_number = offset + 2
        cells = [_cell_to_str(v) for v in row]
        if not any(cells):
            continue  # blank line
        cells += [""] * (len(headers) - len(cells))
        row_warnings: list[str] = []

        def get(name: str) -> str:
            values = [cells[i] for i in all_columns.get(name, [])]
            filled = [v for v in values if v]
            if not filled:
                return ""
            compare = choice_name if name in CHOICE_FIELDS else str
            if len({compare(v).lower() for v in filled}) > 1:
                row_warnings.append(f"{headers[all_columns[name][0]]} has different values "
                                    f"in its repeated columns; using {compare(filled[0])!r}")
            return filled[0]

        raw = {}
        for h, v in zip(headers, cells):
            if h and not raw.get(h):
                raw[h] = v

        url = normalize_linkedin_url(get("url"))
        contact = Contact(
            row_number=row_number,
            url=url,
            first_name=get("first_name"),
            last_name=get("last_name"),
            title=get("title"),
            company=get("company"),
            email=get("email"),
            phone=get("phone"),
            location=get("location"),
            twitter=get("twitter"),
            teal_id=get("teal_id"),
            raw=raw,
        )
        for name in CHOICE_FIELDS:
            setattr(contact, name, choice_name(get(name)))
        for name in DATE_FIELDS:
            value = get(name)
            try:
                parsed = parse_date(value)
            except ValueError:
                row_warnings.append(f"{name.replace('_', ' ')} {value!r} isn't a date; ignoring it")
                parsed = None
            setattr(contact, name, parsed.isoformat() if parsed else "")
        contact.warnings = row_warnings

        if not url:
            contact.problems.append("missing LinkedIn URL")
        elif not LINKEDIN_PROFILE_RE.match(url):
            contact.problems.append(f"not a LinkedIn profile URL: {get('url')}")
        elif url in seen_urls:
            contact.problems.append(f"duplicate of row {seen_urls[url]}")
        else:
            seen_urls[url] = row_number
        if not (contact.first_name or contact.last_name):
            contact.problems.append("missing first and last name")
        contacts.append(contact)

    return Sheet(path, actual_sheet, headers, column_map, contacts, warnings)


def _with_lock_retry(action, path: Path, attempts: int = 4, wait: float = 1.5):
    # OneDrive and Excel hold the file briefly while syncing/saving.
    for attempt in range(attempts):
        try:
            return action()
        except PermissionError as exc:
            if attempt == attempts - 1:
                raise SpreadsheetLockedError(
                    f"Can't write to {path.name} - is it open in Excel (or still syncing)? "
                    "Close it and re-run."
                ) from exc
            time.sleep(wait)


def ensure_columns(sheet: Sheet, fields: list[str]) -> list[str]:
    """Add the tool's own columns (e.g. teal_contact_id) to the header row if
    they're missing. Returns the headers added."""
    missing = [f for f in fields if f not in sheet.column_map]
    if not missing:
        return []
    start = len(sheet.headers)
    while start > 0 and not sheet.headers[start - 1]:
        start -= 1  # reuse trailing blank header cells
    new_headers = {start + i: TOOL_COLUMNS[f] for i, f in enumerate(missing)}

    def add():
        if sheet.path.suffix.lower() == ".csv":
            rows, encoding = _read_csv(sheet.path)
            header = rows[0] + [""] * (start + len(missing) - len(rows[0]))
            for index, name in new_headers.items():
                header[index] = name
            rows[0] = header
            with sheet.path.open("w", newline="", encoding=encoding) as f:
                csv.writer(f).writerows(rows)
        else:
            _write_xlsx_cells(sheet.path, sheet.sheet_name, 1, new_headers)

    _with_lock_retry(add, sheet.path)
    for index, name in new_headers.items():
        if index >= len(sheet.headers):
            sheet.headers += [""] * (index + 1 - len(sheet.headers))
        sheet.headers[index] = name
    for i, f in enumerate(missing):
        sheet.column_map[f] = start + i
    return list(new_headers.values())


def write_back(sheet: Sheet, row_number: int, values: dict[str, str]) -> None:
    """Write canonical fields (e.g. {"email": ..., "phone": ...}) into a row.

    Columns that don't exist in the sheet are skipped. Raises
    SpreadsheetLockedError if the file stays locked (open in Excel, syncing).
    """
    updates = {sheet.column_map[k]: v for k, v in values.items() if k in sheet.column_map}
    if not updates:
        return
    if sheet.path.suffix.lower() == ".csv":
        _with_lock_retry(lambda: _write_back_csv(sheet.path, row_number, updates), sheet.path)
    else:
        _with_lock_retry(lambda: _write_xlsx_cells(sheet.path, sheet.sheet_name, row_number, updates),
                         sheet.path)


def _write_xlsx_cells(path: Path, sheet_name: str | None, row_number: int, updates: dict[int, str]) -> None:
    import openpyxl

    wb = openpyxl.load_workbook(path)
    ws = wb[sheet_name] if sheet_name else wb.worksheets[0]
    for col_index, value in updates.items():
        ws.cell(row=row_number, column=col_index + 1, value=value)
    wb.save(path)


def _write_back_csv(path: Path, row_number: int, updates: dict[int, str]) -> None:
    rows, encoding = _read_csv(path)
    row = rows[row_number - 1]
    for col_index, value in updates.items():
        row += [""] * (col_index + 1 - len(row))
        row[col_index] = value
    with path.open("w", newline="", encoding=encoding) as f:
        csv.writer(f).writerows(rows)
